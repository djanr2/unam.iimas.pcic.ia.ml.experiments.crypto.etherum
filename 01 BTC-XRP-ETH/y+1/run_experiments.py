"""Lanza las corridas de una hoja de experimentos-parametros.xlsx contra el backend.

Uso:  python run_experiments.py [hoja] [tamano_lote]
      python run_experiments.py Sweep_100 10

- Lanza `tamano_lote` corridas (POST /api/tests) y espera a que TODAS terminen antes del
  siguiente lote (el backend ejecuta 4 a la vez; el resto queda en cola).
- Es reanudable: se salta las filas que ya tienen idTest.
- Escribe idTest, RMS_train y RMS_val en el Excel tras cada lote (cierra el Excel antes).
  RMS_train/RMS_val estan en escala NORMALIZADA (0-1), igual que tuple.value_*; RMS_val = fitness.
- Solo lee la base de datos; nunca escribe en ella.
"""
import json, math, sys, time, urllib.request, urllib.error
import psycopg
from openpyxl import load_workbook

XLSX = "experimentos-parametros.xlsx"
API = "http://localhost:8080/api/tests"
DSN = "host=localhost port=5432 dbname=ml_multivariate_db user=postgres password=12345"
POLL_S, MAX_WAIT_S = 20, 3 * 3600
FIELDS = ["idDataset", "seedEga", "seedFaa", "seedDataset", "quasiMinmax", "trainingRate", "numTerm",
          "powerL", "numGeneration", "numIndividual", "crossRate", "mutationRate", "regularizationFactor"]

sheet = sys.argv[1] if len(sys.argv) > 1 else "Sweep_100"
batch_size = int(sys.argv[2]) if len(sys.argv) > 2 else 10


def log(*a):
    print(time.strftime("%H:%M:%S"), *a, flush=True)


def post(payload):
    req = urllib.request.Request(API, data=json.dumps(payload).encode(),
                                 headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, json.loads(r.read())
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode(errors="replace")[:500]


def metrics(cur, id_test):
    """RMS normalizado de entrenamiento y validacion del mejor individuo."""
    cur.execute("select best_individual from test where id_test=%s", (id_test,))
    best = cur.fetchone()[0]
    cur.execute("select fitness from individual where id_individual=%s", (best,))
    fitness = float(cur.fetchone()[0])
    cur.execute("select tp.is_training, tp.value_sample, tp.value_test from tuple tp where tp.id_test=%s", (id_test,))
    acc = {0: [], 1: []}
    for flag, a, b in cur.fetchall():
        acc[flag].append((float(a) - float(b)) ** 2)
    rms = lambda v: math.sqrt(sum(v) / len(v)) if v else None
    return rms(acc[1]), fitness   # (train, val)


def save(wb):
    while True:
        try:
            wb.save(XLSX); return
        except PermissionError:
            log("No puedo guardar el Excel: cierralo (esperando 30 s)..."); time.sleep(30)


wb = load_workbook(XLSX)
ws = wb[sheet]
header = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(header)}
rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col["idTest"]).value in (None, "")]
# filas ya lanzadas pero sin resultados (script interrumpido): esperar y completar
orphans = {r: ws.cell(r, col["idTest"]).value for r in range(2, ws.max_row + 1)
           if ws.cell(r, col["idTest"]).value not in (None, "") and ws.cell(r, col["RMS_val"]).value in (None, "")}
log(f"Hoja {sheet}: {len(rows)} corridas pendientes, lotes de {batch_size}")

db = psycopg.connect(DSN)
db.read_only = True
db.autocommit = True
cur = db.cursor()

def wait_and_fill(ids):
    t0 = time.time(); pending = set(ids.values())
    while pending:
        if time.time() - t0 > MAX_WAIT_S:
            log(f"Tiempo maximo excedido; pendientes: {sorted(pending)}. Abortando."); sys.exit(2)
        cur.execute("select id_test from test where id_test = any(%s) and best_individual is not null", (list(pending),))
        pending -= {x[0] for x in cur.fetchall()}
        log(f"  terminadas {len(ids) - len(pending)}/{len(ids)}")
        if pending: time.sleep(POLL_S)
    for r, tid in ids.items():
        tr, va = metrics(cur, tid)
        ws.cell(r, col["RMS_train"]).value = tr
        ws.cell(r, col["RMS_val"]).value = va
    save(wb)


if orphans:
    log(f"Completando {len(orphans)} corridas ya lanzadas: idTest {sorted(orphans.values())}")
    wait_and_fill(orphans)

for b in range(0, len(rows), batch_size):
    chunk = rows[b:b + batch_size]
    ids = {}
    for r in chunk:
        payload = {k: ws.cell(r, col[k]).value for k in FIELDS}
        status, resp = post(payload)
        if status != 201:
            log(f"ERROR {status} en la fila {r} (run {ws.cell(r, 1).value}): {resp}. Abortando.")
            save(wb); sys.exit(1)
        ids[r] = resp["idTest"]
        ws.cell(r, col["idTest"]).value = ids[r]
        if b == 0 and r == chunk[0]:
            # verificacion de la primera corrida: la base debe guardar los parametros enviados
            cur.execute("select id_dataset,seed_ega,seed_faa,num_term,power_l,num_generation,num_individual,"
                        "cross_rate,mutation_rate from test where id_test=%s", (ids[r],))
            got = cur.fetchone()
            exp = (payload["idDataset"], payload["seedEga"], payload["seedFaa"], payload["numTerm"], payload["powerL"],
                   payload["numGeneration"], payload["numIndividual"], payload["crossRate"], payload["mutationRate"])
            ok = all(abs(float(x) - float(y)) < 1e-9 for x, y in zip(got, exp))
            log("Verificacion 1a corrida (parametros en BD == enviados):", "OK" if ok else f"DISTINTO {got} vs {exp}")
            if not ok:
                save(wb); sys.exit(1)
    save(wb)
    log(f"Lote {b // batch_size + 1}: lanzadas {len(ids)} (idTest {min(ids.values())}-{max(ids.values())})")

    t0 = time.time()
    pending = set(ids.values())
    while pending:
        if time.time() - t0 > MAX_WAIT_S:
            log(f"Tiempo maximo excedido; pendientes: {sorted(pending)}. Abortando."); sys.exit(2)
        time.sleep(POLL_S)
        cur.execute("select id_test from test where id_test = any(%s) and best_individual is not null",
                    (list(pending),))
        pending -= {x[0] for x in cur.fetchall()}
        log(f"  terminadas {len(ids) - len(pending)}/{len(ids)}")

    for r, tid in ids.items():
        tr, va = metrics(cur, tid)
        ws.cell(r, col["RMS_train"]).value = tr
        ws.cell(r, col["RMS_val"]).value = va
    save(wb)
    log(f"Lote {b // batch_size + 1} completo y guardado en el Excel")

log("Todas las corridas terminadas")
